#!/usr/bin/env python3
"""Convertit le fichier annuel de la DRH (.xlsx) en CSV normalisé pour l'import ISS.

    python3 scripts/drh_xlsx_to_csv.py "CNPS DRH 2026.xlsx" data/drh-2026.csv
    python3 scripts/drh_xlsx_to_csv.py "CNPS DRH 2026.xlsx" data/drh-2026.csv.gz   # compressé

Les lignes saisies deux fois (meme matricule et meme contenu) sont supprimees.
Un meme matricule porte par des lignes differentes est conserve et signale :
l'arbitrage appartient a la DRH. L'option --doublons <fichier.csv> ecrit la liste
de ces matricules a lui renvoyer — ce fichier contient des identifiants, il se
manipule comme le fichier source.

Ne conserve que ce dont la carte sanitaire a besoin : aucune donnée identifiante
(ni matricule, ni nom, ni date de naissance exacte — seule l'année est gardée).
Le format produit est décrit dans docs/drh-format.md.

Une destination en .gz est ecrite compressee : le fichier annuel fait ~1,4 Mo,
au-dessus de la limite d'envoi par defaut des reverse proxies, et ~70 Ko gzippe.
"""
import csv, gzip, re, sys, unicodedata
from datetime import datetime

COLONNES = ['region','prefecture','sous_prefecture','structure_affectation','structure_rattachement',
            'profession','profession_oms','hierarchie','statut','sexe','annee_naissance','zone','niveau_structure']

# Matricules de remplissage : ce ne sont pas des identifiants, plusieurs agents
# distincts les partagent. Ne jamais dedoublonner dessus.
MATRICULES_FACTICES = {'', 'nd', 'n d', 'na', 'n a', '0', '-', 'neant', 'aucun'}

# En-têtes acceptés dans le fichier DRH → colonne normalisée (insensible à la casse et aux accents)
SOURCE = {
    'region':'region', 'prefecture commune':'prefecture', 'sous prefecture':'sous_prefecture',
    'structure affectation':'structure_affectation', 'structure':'structure_rattachement',
    'profession':'profession', 'profession cnps oms':'profession_oms', 'hierarchie':'hierarchie',
    'statut employe':'statut', 'sexe':'sexe', 'date de naissance':'annee_naissance',
    'zone':'zone', 'niveau structure':'niveau_structure',
}

def norm(s):
    s = unicodedata.normalize('NFD', str(s or '').lower())
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()

def clean_sexe(v):
    k = norm(v)
    return 'F' if k.startswith('f') else ('H' if k.startswith('h') or k.startswith('m') else '')

def clean_zone(v):
    k = norm(v)
    return 'rurale' if 'rural' in k else ('urbaine' if 'urbain' in k else '')

def clean_niveau(v):
    k = norm(v)
    for mot, out in (('tertiaire','tertiaire'), ('secondaire','secondaire'), ('primaire','primaire')):
        if mot in k:
            return out
    return ''

def annee(v):
    if isinstance(v, datetime):
        return v.year
    m = re.search(r'(19|20)\d{2}', str(v or ''))
    return int(m.group(0)) if m else ''

def main(src, dst, rapport_doublons=None):
    import openpyxl
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb['BASE'] if 'BASE' in wb.sheetnames else wb[wb.sheetnames[0]]

    header, idx = None, {}
    rows_out, ignored = [], 0
    mat_col = None
    vus = {}                 # signature complete -> matricule, pour les doublons stricts
    doublons_stricts = 0
    ambigus = {}             # matricule -> nb de lignes, memes matricules mais lignes differentes
    for raw in ws.iter_rows(values_only=True):
        if header is None:
            # la ligne d'en-tête est la première qui contient « Profession » et « Région »
            vals = [norm(c) for c in raw]
            if 'profession' in vals and 'region' in vals:
                header = vals
                for i, h in enumerate(header):
                    if h in SOURCE:
                        idx[SOURCE[h]] = i
                    if 'matricule' in h:
                        mat_col = i
                manquantes = [c for c in ('region','prefecture','profession','sexe') if c not in idx]
                if manquantes:
                    sys.exit(f"Colonnes obligatoires absentes du fichier : {', '.join(manquantes)}")
            continue
        if not any(v not in (None, '') for v in raw):
            continue
        get = lambda c: (str(raw[idx[c]]).strip() if c in idx and raw[idx[c]] not in (None, '') else '')
        rec = {c: get(c) for c in COLONNES}
        rec['sexe'] = clean_sexe(get('sexe'))
        rec['zone'] = clean_zone(get('zone'))
        rec['niveau_structure'] = clean_niveau(get('niveau_structure'))
        rec['annee_naissance'] = annee(raw[idx['annee_naissance']] if 'annee_naissance' in idx else '')
        if not rec['region'] and not rec['prefecture'] and not rec['profession']:
            ignored += 1
            continue

        # Dedoublonnage : le fichier 2026 contient des lignes saisies deux fois.
        # On ne supprime que ce qui est certain — meme matricule ET meme ligne a
        # l'identique, date de naissance complete comprise. Un meme matricule
        # porte par des lignes qui different (structure, date, profession) peut
        # etre une erreur de saisie comme deux personnes distinctes : on le
        # conserve et on le signale, l'arbitrage appartient a la DRH.
        mat = str(raw[mat_col]).strip() if mat_col is not None and raw[mat_col] not in (None, '') else ''
        if norm(mat) not in MATRICULES_FACTICES:
            # La signature porte sur les donnees, pas sur la ligne brute : le
            # fichier contient une colonne de numerotation qui rendrait chaque
            # ligne unique. Date de naissance complete incluse, plus fine que
            # l'annee conservee en sortie.
            naissance = str(raw[idx['annee_naissance']]) if 'annee_naissance' in idx else ''
            signature = (mat, naissance) + tuple(rec[c] for c in COLONNES)
            if signature in vus:
                doublons_stricts += 1
                continue
            vus[signature] = mat
            ambigus[mat] = ambigus.get(mat, 0) + 1

        rows_out.append(rec)

    opener = gzip.open if dst.endswith('.gz') else open
    with opener(dst, 'wt', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=COLONNES, delimiter=';')
        w.writeheader()
        w.writerows(rows_out)

    n_ambigus = sum(1 for n in ambigus.values() if n > 1)
    sans_annee = sum(1 for r in rows_out if not r['annee_naissance'])
    sans_struct = sum(1 for r in rows_out if not r['structure_affectation'] and not r['structure_rattachement'])
    print(f"{len(rows_out)} agents écrits dans {dst}")
    print(f"  colonnes trouvées : {len(idx)}/{len(COLONNES)}")
    print(f"  sans année de naissance : {sans_annee}   sans libellé de structure : {sans_struct}   lignes vides ignorées : {ignored}")
    print(f"  doublons stricts supprimés : {doublons_stricts}   matricules en double avec des différences (conservés) : {n_ambigus}")
    if rapport_doublons and n_ambigus:
        with open(rapport_doublons, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f, delimiter=';')
            w.writerow(['matricule', 'n_lignes'])
            for mat, n in sorted(ambigus.items(), key=lambda kv: -kv[1]):
                if n > 1:
                    w.writerow([mat, n])
        print(f"  liste des matricules à arbitrer écrite dans {rapport_doublons} (contient des identifiants)")
    print("  aucune donnée identifiante conservée (ni matricule, ni nom, ni date de naissance exacte)")

if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    rapport = None
    if '--doublons' in sys.argv:
        i = sys.argv.index('--doublons')
        if i + 1 >= len(sys.argv):
            sys.exit('--doublons attend un nom de fichier')
        rapport = sys.argv[i + 1]
        args = [a for a in args if a != rapport]
    if len(args) != 2:
        sys.exit(__doc__)
    main(args[0], args[1], rapport)
